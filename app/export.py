# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (c) 2026 Nrupal Akolkar
"""Export chat results to MD, DOCX, HTML, CSV (v0.10.0)."""

from __future__ import annotations

import csv
import html
import io
import re

# Markdown table detection: | col | col | followed by |---|---|
_TABLE_RE = re.compile(
    r"^\s*\|(.+)\|\s*$\n^\s*\|[\s\-:|]+\|\s*$(?:\n^\s*\|(.+)\|\s*$)+",
    re.MULTILINE,
)


def to_markdown(messages: list[dict]) -> str:
    """Convert chat messages to a Markdown document."""
    lines = ["# Milo Export", ""]
    for m in messages:
        role = m.get("role", "")
        content = m.get("content", "")
        if role == "user":
            lines.append(f"## You\n\n{content}\n")
        elif role == "assistant":
            lines.append(f"## Milo\n\n{content}\n")
    return "\n".join(lines)


def to_html(messages: list[dict]) -> str:
    """Convert chat messages to a styled HTML document."""
    parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        "<title>Milo Export</title>",
        "<style>body{font-family:system-ui,sans-serif;max-width:720px;"
        "margin:2em auto;padding:0 1em;line-height:1.6}"
        ".user{background:#f0f4ff;border-radius:8px;padding:1em;margin:1em 0}"
        ".assistant{background:#f8f8f8;border-radius:8px;padding:1em;margin:1em 0}"
        "table{border-collapse:collapse;width:100%;margin:1em 0}"
        "th,td{border:1px solid #ddd;padding:8px;text-align:left}"
        "th{background:#f0f0f0}pre{background:#f5f5f5;padding:1em;overflow-x:auto;border-radius:4px}</style>",
        "</head><body><h1>Milo Export</h1>",
    ]
    for m in messages:
        role = m.get("role", "")
        content = html.escape(m.get("content", ""))
        # Simple markdown-ish rendering: preserve line breaks.
        content = content.replace("\n", "<br>")
        cls = "user" if role == "user" else "assistant"
        label = "You" if role == "user" else "Milo"
        parts.append(f"<div class='{cls}'><strong>{label}</strong><br>{content}</div>")
    parts.append("</body></html>")
    return "".join(parts)


def to_docx(messages: list[dict]) -> bytes:
    """Convert chat messages to a Word .docx file."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Milo Export", level=1)
    for m in messages:
        role = m.get("role", "")
        content = m.get("content", "")
        label = "You" if role == "user" else "Milo"
        doc.add_heading(label, level=2)
        # Split on double newlines for paragraphs.
        for para in content.split("\n\n"):
            para = para.strip()
            if para:
                p = doc.add_paragraph(para)
                p.style.font.size = Pt(11)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def extract_tables(markdown_text: str) -> list[list[list[str]]]:
    """Extract markdown tables as lists of rows."""
    tables = []
    for match in _TABLE_RE.finditer(markdown_text):
        block = match.group(0)
        rows = []
        for line in block.strip().split("\n"):
            line = line.strip()
            if re.match(r"^\|[\s\-:|]+\|$", line):
                continue  # separator row
            cells = [c.strip() for c in line.strip("|").split("|")]
            rows.append(cells)
        if rows:
            tables.append(rows)
    return tables


def tables_to_csv(tables: list[list[list[str]]]) -> str:
    """Convert extracted tables to CSV (multiple tables separated)."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    for i, table in enumerate(tables):
        if i > 0:
            writer.writerow([])  # blank row between tables
        for row in table:
            writer.writerow(row)
    return buf.getvalue()
